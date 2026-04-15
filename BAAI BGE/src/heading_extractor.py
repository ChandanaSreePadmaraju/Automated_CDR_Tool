"""
heading_extractor.py
--------------------
Extract all heading paragraphs from a .docx file.
Returns a flat list so both template and data docs can be compared.
"""

from docx import Document


def extract_headings(doc_path: str) -> list[dict]:
    """
    Read every paragraph in *doc_path* and return those whose style name
    starts with "Heading".

    Each entry in the returned list is:
    {
        "text":            str  – heading text (stripped),
        "level":           int  – heading level (1, 2, 3 …),
        "paragraph_index": int  – position in doc.paragraphs (0-based),
    }

    Empty heading paragraphs are skipped.
    """
    doc = Document(doc_path)
    headings: list[dict] = []

    for idx, para in enumerate(doc.paragraphs):
        style_name = para.style.name          # e.g. "Heading 1", "Heading 2"
        if not style_name.startswith("Heading"):
            continue

        text = para.text.strip()
        if not text:
            continue                          # skip blank headings

        parts = style_name.split()
        level = int(parts[-1]) if parts[-1].isdigit() else 1

        headings.append({
            "text":            text,
            "level":           level,
            "paragraph_index": idx,
        })

    return headings
